"""Write the video-ideas table into README.md, between the IDEAS markers.

Reads video-ideas.xlsx (sheet "Ideas"). The id is computed from the row number, because
the id column is a formula and files saved by scripts have no stored formula results.
The title links to videos/NN-*.md when that file exists (NN = the id, two digits).
The last column shows my_take, or notes when there is no take yet, shortened.
"""
import datetime
import glob
import os
import re
import sys
from collections import Counter

from openpyxl import load_workbook

XLSX = "video-ideas.xlsx"
README = "README.md"
START, END = "<!-- IDEAS:START -->", "<!-- IDEAS:END -->"
HEADERS = ["id", "Working title", "Topic", "Status", "Script", "Priority", "Added", "Take / notes"]
TAKE_LEN = 140
STATUS_ORDER = ["riff pending", "riff recorded", "scripting", "recording", "animating", "rendered", "published", "parked"]


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    t = str(v).replace("\r", " ").replace("\n", " ").strip()
    t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("|", "\\|")
    return t


def video_link(idnum, title):
    hits = sorted(glob.glob(os.path.join("videos", f"{int(idnum):02d}-*.md")))
    return f"[{title}]({hits[0].replace(os.sep, '/')})" if hits else title


def shorten(v):
    t = " ".join(str(v).split())
    return t if len(t) <= TAKE_LEN else t[: TAKE_LEN - 1].rstrip() + "\u2026"


def main():
    ws = load_workbook(XLSX)["Ideas"]
    col = {c.value: c.column for c in ws[1] if c.value}
    get = lambda r, k: ws.cell(row=r, column=col[k]).value
    rows = []
    for r in range(2, ws.max_row + 1):
        if not get(r, "working_title"):
            continue
        idnum = f"{r - 1:03d}"
        rows.append({
            "status": get(r, "status"),
            "cells": [
                idnum,
                video_link(idnum, cell_text(get(r, "working_title"))),
                cell_text(get(r, "topic_area")),
                cell_text(get(r, "status")),
                cell_text(get(r, "script_status")),
                cell_text(get(r, "priority")),
                cell_text(get(r, "date_added")),
                cell_text(shorten(get(r, "my_take") or get(r, "notes") or "")),
            ],
        })

    counts = Counter(str(x["status"]) for x in rows)
    summary = ", ".join(f"{counts[s]} {s}" for s in STATUS_ORDER if counts[s]) or "no ideas yet"
    lines = [f"{len(rows)} idea{'s' if len(rows) != 1 else ''}: {summary}.", ""]
    lines.append("| " + " | ".join(HEADERS) + " |")
    lines.append("|" + "|".join("---" for _ in HEADERS) + "|")
    for x in rows:
        lines.append("| " + " | ".join(x["cells"]) + " |")
    block = "\n".join(lines)

    text = open(README, encoding="utf-8").read()
    pat = re.compile(re.escape(START) + r".*?" + re.escape(END), re.S)
    if not pat.search(text):
        print(f"README.md is missing the {START} / {END} markers.", file=sys.stderr)
        sys.exit(1)
    new = pat.sub(lambda m: f"{START}\n{block}\n{END}", text)
    if new != text:
        open(README, "w", encoding="utf-8", newline="\n").write(new)
        print("README updated.")
    else:
        print("README already up to date.")


if __name__ == "__main__":
    main()
