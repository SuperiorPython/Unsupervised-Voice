"""Append one video idea, taken from the GitHub issue form, to video-ideas.xlsx.

Reads the issue body from the ISSUE_BODY environment variable. Writes the new row's
number to $GITHUB_OUTPUT (row, id). Exits non-zero with a plain message when the idea
can't be added, and the workflow comments that message on the issue.
"""
import datetime
import os
import re
import sys

from openpyxl import load_workbook

XLSX = "video-ideas.xlsx"
LAST_ROW = 300  # the sheet's dropdowns and ids are pre-set down to this row

# issue-form label -> spreadsheet column header
LABEL_TO_HEADER = {
    "Working title": "working_title",
    "Topic area": "topic_area",
    "Status": "status",
    "Origin": "origin",
    "My take": "my_take",
    "Best moments": "best_moments",
    "Hook idea": "hook_idea",
    "Riff recorded": "riff_recorded",
    "Script status": "script_status",
    "Priority": "priority",
    "Notes": "notes",
}
# These columns only accept a value from the Lists sheet (same as the dropdowns in Excel).
STRICT = ("status", "riff_recorded", "script_status", "priority")


def fail(msg):
    print(msg, file=sys.stderr)
    sys.exit(1)


def parse(body):
    """Issue forms render as '### Label' followed by the answer."""
    parts = re.split(r"^### +(.+?) *$", body, flags=re.M)
    out = {}
    for label, text in zip(parts[1::2], parts[2::2]):
        t = text.strip()
        out[label.strip()] = "" if t == "_No response_" else t
    return out


def main():
    fields = parse(os.environ.get("ISSUE_BODY", ""))
    values = {LABEL_TO_HEADER[k]: v for k, v in fields.items() if k in LABEL_TO_HEADER}
    if not values.get("working_title"):
        fail("The form has no working title.")

    wb = load_workbook(XLSX)
    ws = wb["Ideas"]
    lists = wb["Lists"]
    col = {c.value: c.column for c in ws[1] if c.value}

    allowed = {}
    for c in lists[1]:
        if not c.value:
            continue
        vals, r = [], 2
        while lists.cell(row=r, column=c.column).value is not None:
            vals.append(str(lists.cell(row=r, column=c.column).value))
            r += 1
        allowed[c.value] = vals
    for h in STRICT:
        if values.get(h) and values[h] not in allowed.get(h, []):
            fail(f"'{values[h]}' isn't an allowed {h}. Allowed: {', '.join(allowed.get(h, []))}. "
                 f"Update the issue form and the Lists sheet so they match.")

    row = next((r for r in range(2, LAST_ROW + 1) if not ws.cell(row=r, column=col["working_title"]).value), None)
    if row is None:
        fail(f"The tracker is full (rows 2-{LAST_ROW}).")

    for h, v in values.items():
        cell = ws.cell(row=row, column=col[h])
        if h == "priority" and v:
            cell.value = int(v)
        else:
            cell.value = v
            if isinstance(v, str):
                cell.data_type = "s"  # always plain text, even if it starts with "="
    ws.cell(row=row, column=col["date_added"]).value = datetime.date.today()
    wb.save(XLSX)

    idnum = f"{row - 1:03d}"
    print(f"Added row {row} (id {idnum}): {values['working_title']}")
    out = os.environ.get("GITHUB_OUTPUT")
    if out:
        with open(out, "a", encoding="utf-8") as f:
            f.write(f"row={row}\nid={idnum}\n")


if __name__ == "__main__":
    main()
